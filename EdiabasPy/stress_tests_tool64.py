# -*- coding: utf-8 -*-
"""
Created on Mon Jun 17 14:12:09 2024

@author: testhouse
"""
import sys
sys.path.insert(0, r'E:\Scripts\dlt_tools')
sys.path.insert(0, r'E:\Scripts\umlauttestrack')
import clr
from openpyxl import load_workbook  # You need openpyxl to handle existing Excel files
import pandas as pd
from datetime import datetime
import re
import time
from dltmodules.dlt_export import DltExport
from dltmodules import connector
#from dltmodules import dlt_filter
#import subprocess
from webcam import WebcamRecorder  # Import the WebcamRecorder class
import threading
import wakeUp
import clampOnOff

# Add reference to the C# DLL and import the necessary class
clr.AddReference(r"E:\Scripts\EdiabasPy\DiagnosticModuleDLL\DiagnosticModule.dll")
from DiagnosticModule import EdiabasSingleton

# Instantiate the C# class. This is used to access the Tool32/64 API
csharp_instance = EdiabasSingleton()

#ip_address_IDC = "169.254.157.47"      #IDC
ip_address_IDC = "169.254.166.99"         #IDCevo
ip_address = "169.254.61.64"              #BCP
screen = True
amount = 1000000000
csv = ''
base_filename = r'E:\DLT_Traces\IDCevo\loop'
excel_filename = r'E:\DLT_Traces\IDCevo\error_counts.xlsx'
python = r"C:\Users\testhouse\anaconda3\envs\ecutest\python.exe"
base_command = [python, "E:/Scripts/EdiabasPy/SendUDP/sendhsfz_2.py", "--ip-addr", ip_address]

video_filename = r'E:\Videos\IDCevo\output_3.avi'
base_filename_pic = r'E:\TestPics\IDCevo\loop'
picture_filename = 'picture.jpg'
hu_sgbd = "IDCEVO25"
ipb_sgbd = "IPB_APP1"
fs_directory = r"E:\DTC_Reports"


# Define a list of tuples with the variable parts of the command, durations, and messages
# For example: (command_parts, duration_in_seconds, "Message to print")
commands_with_duration_messages = [
    (["PAD"], 80, "PAD sent"),
    (["wohnen"], 10, "wohnen sent"),
    (["parken"], 60, "parken sent"),
    (["wohnen"], 10, "wohnen sent"),
    # Add more commands, durations, and messages here
]

# Start logging dlt with variables provided
def start_logging(filename):
    exportOptions = DltExport(None, screen, csv, filename, amount)
    c = connector.DltConnectorThread(ip_address_IDC, exp=exportOptions, buffering=False)
    c.start_reading()
    return c

# Stop the logging and save
def stop_logging(c):
    c.stop_reading()

# Function to count errors in DLT file
def count_errors_in_dlt_file(file_path):
    error_pattern = re.compile(r'DEL NEIGHBOUR:INCOMPLETE|NEW NEIGHBOUR:FAILED')
    error_count = 0
    try:
        with open(file_path, 'rb') as file:
            for line in file:
                decoded_line = line.decode('latin-1')
                if error_pattern.search(decoded_line):
                    error_count += 1
    except FileNotFoundError:
        print(f"File {file_path} not found.")
        return None
    except Exception as e:
        print(f"An error occurred while reading the file: {e}")
        return None
    return error_count

# Function to append row to Excel
def append_row_to_excel(filepath, row):
    try:
        # Try to read the existing Excel file
        df = pd.read_excel(filepath)
    except FileNotFoundError:
        # If the file does not exist, create a new DataFrame
        df = pd.DataFrame(columns=['Filename', 'Timestamp', 'Number of Errors'])
    # Convert the row to a DataFrame and append it
    row_df = pd.DataFrame([row])  # Convert row to a DataFrame with a single row
    df = pd.concat([df, row_df], ignore_index=True)  # Use concat instead of append
    df.to_excel(filepath, index=False)
    
def append_row_to_excel_sheet(filepath, row, sheet_name='Sheet1'):
    # Create a DataFrame from the row
    row_df = pd.DataFrame([row])

    try:
        # Try to open the existing workbook
        book = load_workbook(filepath)
        # If the sheet exists, load it, else create a new sheet
        if sheet_name in book.sheetnames:
            writer = pd.ExcelWriter(filepath, engine='openpyxl')  # Specify the ExcelWriter
            writer.book = book  # Load the existing workbook
            writer.sheets = {ws.title: ws for ws in book.worksheets}  # Load existing sheets
            # Get the last row in the existing Excel sheet
            # If the sheet is not empty, start at the first empty row after the existing content.
            startrow = writer.sheets[sheet_name].max_row
            # Append without writing index
            row_df.to_excel(writer, sheet_name=sheet_name, startrow=startrow, index=False, header=False)
        else:
            # The sheet does not exist, we can just write to it as a new sheet
            row_df.to_excel(filepath, sheet_name=sheet_name, index=False)
        writer.save()
        writer.close()
    except FileNotFoundError:
        # If the file does not exist, simply use DataFrame.to_excel to create it
        row_df.to_excel(filepath, sheet_name=sheet_name, index=False)


recorder = WebcamRecorder(video_filename, picture_filename)

# Start recording before the while loop
threading.Thread(target=recorder.start_recording).start()

try:
    iteration = 0
    while iteration < 300:
        iteration += 1
        if (iteration % 10 == 0):
            clampOnOff.set_status_all("off")
            time.sleep(2)
            clampOnOff.set_status_all("on")
            time.sleep(40)
            
        current_filename = f"{base_filename}_{iteration}.dlt"
        try:
            c = start_logging(current_filename)
        except Exception as e:
            print(f"Error occurred in DLT logging start: {e}")
        
        
        for command, duration, message in commands_with_duration_messages:
            print("Sending command:", command)
            if command[0] == "PAD":
                csharp_instance.Job("IPB_APP1","steuern_routine","ID;0x1031;STR;0x07")
            elif command[0] == "wohnen":
                csharp_instance.Job("IPB_APP1","steuern_routine","ID;0x1031;STR;0x05")
            elif command[0] == "parken":
                csharp_instance.Job("IPB_APP1","steuern_routine","ID;0x1031;STR;0x01")
            print(message)
            
            if message == "wohnen sent":
                wakeUp.send_wake_up()

            if message == "PAD sent":
                # wakeUp.send_wake_up()
                # Schedule the picture to be taken 19 seconds after "PAD sent"
                picture_filename = f"{base_filename_pic}_{iteration}.jpg"
                threading.Timer(19, recorder.capture_picture, [picture_filename]).start()
            
                try:
                    csharp_instance.Job(hu_sgbd, "fs_lesen_expert", ", ")
                except Exception as e:
                    print(f"Error occurred in csharp_instance.Job: {e}")
                
                saveName = f"loop_{iteration}"
                csharp_instance.Results_1(fs_directory, saveName)

            print(f"Waiting for {duration} seconds...")
            time.sleep(duration) 
        try:
            stop_logging(c)
        except Exception as e:
            print(f"Error occurred in DLT logging stop: {e}")
        stop_logging(c)
        
        error_count = count_errors_in_dlt_file(current_filename)
        if error_count is not None:
            row = {
                'Filename': current_filename,
                'Timestamp': datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                'Number of Errors': error_count
            }
            append_row_to_excel(excel_filename, row)

finally:
    # Stop recording after the while loop
    recorder.stop_recording()

